"""Rendered resumes must be ATS-friendly: real text, no layout tables, all four
templates producing a valid document."""

from __future__ import annotations

import io
import zipfile
from datetime import date

import pytest
from jobapply_resume.models import (
    ContactInfo,
    TailoredBullet,
    TailoredEducation,
    TailoredExperience,
    TailoredResume,
)
from jobapply_resume.renderers import TEMPLATES, DocxRenderer, PdfRenderer, date_range


@pytest.fixture
def document() -> TailoredResume:
    return TailoredResume(
        contact=ContactInfo(
            full_name="Jordan Rivera",
            email="jordan@example.com",
            phone="(415) 555-0142",
            location="San Francisco, CA",
            linkedin_url="https://linkedin.com/in/jordanrivera",
        ),
        summary="Data engineer with six years building batch and streaming pipelines.",
        summary_source_ids=["profile"],
        skills=["Python", "SQL", "Airflow", "dbt"],
        experience=[
            TailoredExperience(
                experience_id="experience_1",
                company="Northwind Analytics",
                title="Senior Data Engineer",
                location="San Francisco, CA",
                start_date=date(2021, 3, 1),
                is_current=True,
                bullets=[
                    TailoredBullet(
                        text="Built automated data pipelines using Python & SQL.",
                        source_ids=["experience_1"],
                        confidence=0.95,
                    )
                ],
            )
        ],
        education=[
            TailoredEducation(
                education_id="education_1",
                institution="University of Texas at Austin",
                degree="B.S.",
                field_of_study="Computer Science",
                end_date=date(2018, 5, 1),
                gpa=3.7,
            )
        ],
        certifications=["AWS Certified Solutions Architect – Associate"],
    )


@pytest.mark.parametrize("template", sorted(TEMPLATES))
def test_every_template_renders_both_formats(document, template):
    document.template = template
    docx_bytes = DocxRenderer(template).render(document)
    pdf_bytes = PdfRenderer(template).render(document)
    assert docx_bytes.startswith(b"PK")
    assert pdf_bytes.startswith(b"%PDF")


def test_pdf_text_is_selectable_and_complete(document):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(PdfRenderer().render(document)))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Jordan Rivera" in text
    assert "jordan@example.com" in text
    assert "Northwind Analytics" in text
    assert "Built automated data pipelines" in text
    assert "University of Texas at Austin" in text
    assert "AWS Certified Solutions Architect" in text


def test_a_short_resume_fits_on_one_page(document):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(PdfRenderer().render(document)))
    assert len(reader.pages) == 1


def test_ampersands_are_escaped_rather_than_breaking_the_pdf(document):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(PdfRenderer().render(document)))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Python & SQL" in text


def test_docx_uses_paragraphs_not_tables(document):
    """Tables are the classic way to break an ATS parser."""
    import docx

    data = DocxRenderer().render(document)
    parsed = docx.Document(io.BytesIO(data))
    assert parsed.tables == []
    text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "Jordan Rivera" in text
    assert "Senior Data Engineer, Northwind Analytics" in text


def test_docx_contains_no_images(document):
    data = DocxRenderer().render(document)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media == []


def test_current_roles_read_as_present(document):
    assert date_range(date(2021, 3, 1), None, True).endswith("Present")
    assert date_range(date(2018, 6, 1), date(2021, 2, 1), False) == "Jun 2018 – Feb 2021"
    assert date_range(None, None, False) == ""


def test_templates_differ_in_typography_not_structure():
    keys = {TEMPLATES[name].body_font for name in TEMPLATES}
    assert len(keys) > 1, "templates should not all look identical"
